"""Turn any uploaded file into text elements, then into overlapping chunks."""
from dataclasses import dataclass
from docx import Document
from io import BytesIO
from pathlib import Path

import pandas as pd

import pdfplumber
from backend import config, llm


@dataclass
class Element:
    text: str
    kind: str          # "text" | "table" | "image"
    page: int = 0      # 0 = not applicable


# helpers
def rows_to_markdown(rows: list[list]) -> str:
    """List-of-rows (from pdfplumber / docx) -> markdown table."""
    rows = [[("" if c is None else str(c).replace("\n", " ").strip()) for c in r] for r in rows if r]
    if len(rows) < 2:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


def dataframe_elements(df: pd.DataFrame, label: str) -> list[Element]:
    df = df.dropna(how="all").dropna(axis=1, how="all").head(5000)
    if df.empty:
        return []
    summary = (
        f"Table '{label}': {len(df)} rows, {len(df.columns)} columns. "
        f"Columns: {', '.join(map(str, df.columns))}."
    )
    return [Element(summary, "text"), Element(df.fillna("").to_markdown(index=False), "table")]


# per-format loaders
def load_pdf(data: bytes) -> list[Element]:

    out: list[Element] = []
    vision_used = 0
    with pdfplumber.open(BytesIO(data)) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            text = (page.extract_text() or "").strip()
            for tbl in page.extract_tables():
                md = rows_to_markdown(tbl)
                if md:
                    out.append(Element(md, "table", i))
            if len(text) >= 30:
                out.append(Element(text, "text", i))
            elif vision_used < config.MAX_VISION_PDF_PAGES:   # scanned page -> vision model
                buf = BytesIO()
                page.to_image(resolution=150).original.save(buf, "PNG")
                out.append(Element(llm.describe_image(buf.getvalue()), "image", i))
                vision_used += 1
    return out


def load_docx(data: bytes) -> list[Element]:

    doc = Document(BytesIO(data))
    out = []
    body = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    if body:
        out.append(Element(body, "text"))
    for t in doc.tables:
        md = rows_to_markdown([[c.text for c in r.cells] for r in t.rows])
        if md:
            out.append(Element(md, "table"))
    return out


def load_csv(data: bytes, name: str) -> list[Element]:
    return dataframe_elements(pd.read_csv(BytesIO(data)), name)


def load_xlsx(data: bytes, name: str) -> list[Element]:
    out = []
    for sheet, df in pd.read_excel(BytesIO(data), sheet_name=None).items():
        out += dataframe_elements(df, f"{name} / {sheet}")
    return out


def load_text(data: bytes) -> list[Element]:
    return [Element(data.decode("utf-8", errors="ignore"), "text")]


def load_image(data: bytes) -> list[Element]:
    return [Element(llm.describe_image(data), "image")]


def load_file(name: str, data: bytes) -> list[Element]:
    ext = Path(name).suffix.lower().lstrip(".")
    if ext == "pdf":
        return load_pdf(data)
    if ext == "docx":
        return load_docx(data)
    if ext == "csv":
        return load_csv(data, name)
    if ext == "xlsx":
        return load_xlsx(data, name)
    if ext in ("txt", "md"):
        return load_text(data)
    if ext in config.IMAGE_TYPES:
        return load_image(data)
    raise ValueError(f"Unsupported file type: .{ext}")


# chunking
def split_text(text: str, size: int, overlap: int) -> list[str]:
    text = text.strip()
    if len(text) <= size:
        return [text] if text else []
    chunks, start = [], 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            cut = max(text.rfind("\n\n", start, end), text.rfind(". ", start, end))
            if cut > start + size // 2:
                end = cut + 1
        chunks.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return [c for c in chunks if c]


def split_table(md: str, rows_per_chunk: int) -> list[str]:
    """Keep the header on every chunk so each piece is self-explanatory."""
    lines = md.splitlines()
    if len(lines) <= 2 + rows_per_chunk:
        return [md]
    header, body = lines[:2], lines[2:]
    return ["\n".join(header + body[i:i + rows_per_chunk]) for i in range(0, len(body), rows_per_chunk)]


def chunk_elements(elements: list[Element], source: str) -> list[dict]:
    chunks = []
    for el in elements:
        if el.kind == "table":
            pieces = split_table(el.text, config.TABLE_ROWS_PER_CHUNK)
        else:
            pieces = split_text(el.text, config.CHUNK_SIZE, config.CHUNK_OVERLAP)
        for p in pieces:
            chunks.append({"text": p, "meta": {"source": source, "page": el.page, "kind": el.kind}})
    return chunks
